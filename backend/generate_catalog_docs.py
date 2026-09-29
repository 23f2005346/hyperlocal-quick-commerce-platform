import sqlite3
import os
import sys
import re
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether, HRFlowable
)
from reportlab.pdfgen import canvas
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

sys.stdout.reconfigure(encoding='utf-8')

ROOT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DB_PATH = os.path.join(ROOT_DIR, 'backend', 'kirana.db')
PDF_OUTPUT = os.path.join(ROOT_DIR, 'Komal_Mart_Product_Checklist.pdf')
XLSX_OUTPUT = os.path.join(ROOT_DIR, 'Komal_Mart_Product_Checklist.xlsx')
DOCX_OUTPUT = os.path.join(ROOT_DIR, 'Komal_Mart_Product_Checklist.docx')

def clean_ascii(text):
    if not text:
        return ""
    # Replace rupee symbol with Rs.
    text = text.replace('₹', 'Rs.')
    # Remove non-latin characters for PDF Helvetica compatibility
    # Keep standard punctuation, numbers, alphabets
    cleaned = re.sub(r'[^\x00-\x7F]+', '', text)
    # Clean double spaces or empty parens
    cleaned = re.sub(r'\(\s*\)', '', cleaned)
    cleaned = re.sub(r'\s+', ' ', cleaned).strip()
    return cleaned

def fetch_data():
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    c.execute('''
        SELECT c.name as category_name, p.id, p.name, p.name_hi, p.is_loose, p.image_url,
               v.unit_size, v.selling_price, v.mrp
        FROM products p
        JOIN categories c ON p.category_id = c.id
        LEFT JOIN product_variants v ON p.id = v.product_id
        ORDER BY c.id, p.id, v.id
    ''')
    rows = c.fetchall()
    
    categories = {}
    for cat_name, pid, pname, pname_hi, is_loose, img_url, unit, price, mrp in rows:
        if cat_name not in categories:
            categories[cat_name] = {}
        if pid not in categories[cat_name]:
            categories[cat_name][pid] = {
                'id': pid,
                'name': pname,
                'name_hi': pname_hi or '',
                'is_loose': is_loose,
                'image_url': img_url or '',
                'variants': []
            }
        if unit:
            categories[cat_name][pid]['variants'].append({
                'unit': unit,
                'price': price,
                'mrp': mrp
            })
    return categories

class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super(NumberedCanvas, self).__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            canvas.Canvas.showPage(self)
        canvas.Canvas.save(self)

    def draw_page_number(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748b"))
        self.setStrokeColor(colors.HexColor("#cbd5e1"))
        self.setLineWidth(0.5)
        self.line(32, 36, 563, 36)
        page_str = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(563, 24, page_str)
        self.drawString(32, 24, "Komal Mart (Wadala, Mumbai) - Master Catalog & Photo Ingestion Checklist")
        self.restoreState()

def generate_pdf(categories):
    doc = SimpleDocTemplate(
        PDF_OUTPUT,
        pagesize=A4,
        leftMargin=30,
        rightMargin=30,
        topMargin=30,
        bottomMargin=46
    )
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=16,
        leading=20,
        textColor=colors.HexColor('#064e3b')
    )
    sub_style = ParagraphStyle(
        'DocSub',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=12,
        textColor=colors.HexColor('#475569')
    )
    cat_header_style = ParagraphStyle(
        'CatHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=15,
        textColor=colors.HexColor('#065f46')
    )
    th_style = ParagraphStyle(
        'TH',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.white
    )
    cell_pname = ParagraphStyle(
        'CellPName',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor('#1e293b')
    )
    cell_var = ParagraphStyle(
        'CellVar',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7,
        leading=9,
        textColor=colors.HexColor('#334155')
    )
    cell_check = ParagraphStyle(
        'CellCheck',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.5,
        leading=9.5,
        alignment=1,
        textColor=colors.HexColor('#0f172a')
    )
    
    story = []
    
    story.append(Paragraph("Komal Mart - Master Store Inventory & Photo Checklist", title_style))
    story.append(Spacer(1, 3))
    story.append(Paragraph("Wadala Storefront | Total Products: 128 | Total Variants: 331 | Use checkboxes below to track saved photos and adjust rates.", sub_style))
    story.append(Spacer(1, 6))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#059669"), spaceAfter=8))

    inst_data = [[
        Paragraph("<b>Checklist Guide:</b><br/>"
                  "1. <b>Photo [ ] Done</b>: Tick when you have downloaded/saved the photo in your folder with price (e.g. <i>toor daal 190 per kg.jpg</i>).<br/>"
                  "2. <b>Del [ ] Cut</b>: Mark if this product is not sold in the shop and should be removed from the catalog.<br/>"
                  "3. <b>New Rate / Notes</b>: Write revised wholesale/retail rates or extra items to add.", sub_style)
    ]]
    inst_table = Table(inst_data, colWidths=[535])
    inst_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#f0fdf4')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#86efac')),
        ('PADDING', (0,0), (-1,-1), 6),
    ]))
    story.append(inst_table)
    story.append(Spacer(1, 10))

    for cat_name, prods in categories.items():
        clean_cat = clean_ascii(cat_name)
        story.append(KeepTogether([
            Paragraph(f"<b>CATEGORY: {clean_cat.upper()}</b> ({len(prods)} Products)", cat_header_style),
            Spacer(1, 3)
        ]))
        
        table_data = [[
            Paragraph("#", th_style),
            Paragraph("Product Name / SKU", th_style),
            Paragraph("Pack Variants & Rate", th_style),
            Paragraph("Photo [ ]", th_style),
            Paragraph("Del [ ]", th_style),
            Paragraph("New Price / Notes", th_style)
        ]]
        
        for idx, (pid, pinfo) in enumerate(prods.items(), start=1):
            var_texts = []
            for v in pinfo['variants']:
                clean_unit = clean_ascii(v['unit'])
                p_val = int(v['price']) if v['price'].is_integer() else v['price']
                var_texts.append(f"- {clean_unit}: <b>Rs.{p_val}</b>")
            var_html = "<br/>".join(var_texts) if var_texts else "Standard"
            
            clean_pname = clean_ascii(pinfo['name'])
            p_title = f"<b>{clean_pname}</b>"
            
            table_data.append([
                Paragraph(str(idx), cell_pname),
                Paragraph(p_title, cell_pname),
                Paragraph(var_html, cell_var),
                Paragraph("[  ] Done", cell_check),
                Paragraph("[  ] Cut", cell_check),
                Paragraph("____________", cell_var)
            ])
            
        t = Table(table_data, colWidths=[20, 185, 145, 55, 48, 82])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#065f46')),
            ('ALIGN', (0,0), (-1,0), 'LEFT'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#f8fafc')]),
            ('PADDING', (0,0), (-1,-1), 3),
            ('BOTTOMPADDING', (0,1), (-1,-1), 4),
            ('TOPPADDING', (0,1), (-1,-1), 4),
        ]))
        story.append(t)
        story.append(Spacer(1, 10))
        
    doc.build(story, canvasmaker=NumberedCanvas)
    print(f"Generated clean PDF: {PDF_OUTPUT}")

def generate_excel(categories):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Komal Mart Master Checklist"
    
    header_fill = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    cat_fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
    cat_font = Font(name="Calibri", size=11, bold=True, color="064E3B")
    bold_font = Font(name="Calibri", size=10, bold=True)
    normal_font = Font(name="Calibri", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    
    headers = [
        "Category", "Product ID", "Product Name (English)", "Regional Name (Hindi/Marathi)",
        "Type", "Pack Variants & Current Rates", "Photo Saved in Folder? [Tick]",
        "Action (Keep / Remove)", "Updated Selling Price / New Rate", "Notes / Revisions"
    ]
    
    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
    
    row_idx = 2
    for cat_name, prods in categories.items():
        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=len(headers))
        cat_cell = ws.cell(row=row_idx, column=1, value=f"CATEGORY: {cat_name.upper()} ({len(prods)} Products)")
        cat_cell.fill = cat_fill
        cat_cell.font = cat_font
        cat_cell.alignment = left_align
        for c in range(1, len(headers) + 1):
            ws.cell(row=row_idx, column=c).border = thin_border
        row_idx += 1
        
        for pid, pinfo in prods.items():
            var_str = ", ".join([f"{v['unit']}: Rs.{int(v['price']) if v['price'].is_integer() else v['price']}" for v in pinfo['variants']])
            row_data = [
                cat_name,
                pid,
                pinfo['name'],
                pinfo['name_hi'],
                "Loose Mandi" if pinfo['is_loose'] else "Packaged FMCG",
                var_str,
                "Pending [  ]",
                "Keep",
                "",
                ""
            ]
            ws.append(row_data)
            for col_num in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col_num)
                cell.font = normal_font
                cell.border = thin_border
                if col_num in [2, 5, 7, 8]:
                    cell.alignment = center_align
                else:
                    cell.alignment = left_align
            row_idx += 1

    col_widths = {1: 22, 2: 12, 3: 38, 4: 30, 5: 16, 6: 44, 7: 24, 8: 20, 9: 30, 10: 30}
    for col_num, width in col_widths.items():
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_num)].width = width
        
    wb.save(XLSX_OUTPUT)
    print(f"Generated Excel: {XLSX_OUTPUT}")

def generate_docx(categories):
    doc = Document()
    
    title = doc.add_paragraph()
    r = title.add_run("Komal Mart — Master Inventory & Photo Checklist")
    r.font.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(6, 78, 59)
    
    sub = doc.add_paragraph()
    r_sub = sub.add_run("Wadala Storefront • 128 Products • 331 Variants • Editable Checklist Document")
    r_sub.font.size = Pt(9.5)
    r_sub.font.color.rgb = RGBColor(100, 116, 139)
    
    doc.add_paragraph("Instructions: Check [x] under 'Photo Saved' when you save the image in your folder (e.g. 'toor daal 190 per kg.jpg'). Mark 'Remove' if the product is not in stock, or write the new price.")
    
    for cat_name, prods in categories.items():
        h = doc.add_heading(level=2)
        r_h = h.add_run(f"{cat_name} ({len(prods)} Products)")
        r_h.font.color.rgb = RGBColor(6, 95, 70)
        
        table = doc.add_table(rows=1, cols=6)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr_cells = table.rows[0].cells
        hdr_titles = ["#", "Product Name", "Pack Variants & Rate", "Photo Saved?", "Action", "New Price / Notes"]
        for i, title_text in enumerate(hdr_titles):
            hdr_cells[i].text = title_text
            for p in hdr_cells[i].paragraphs:
                for run in p.runs:
                    run.font.bold = True
                    run.font.size = Pt(8.5)
                    run.font.color.rgb = RGBColor(255, 255, 255)
        
        for idx, (pid, pinfo) in enumerate(prods.items(), start=1):
            row_cells = table.add_row().cells
            row_cells[0].text = str(idx)
            row_cells[1].text = f"{pinfo['name']} ({pinfo['name_hi']})" if pinfo['name_hi'] else pinfo['name']
            
            var_lines = [f"{v['unit']}: Rs.{int(v['price']) if v['price'].is_integer() else v['price']}" for v in pinfo['variants']]
            row_cells[2].text = "\n".join(var_lines)
            row_cells[3].text = "[   ] Done"
            row_cells[4].text = "[   ] Remove"
            row_cells[5].text = ""
            
            for cell in row_cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(8.5)
                        
        doc.add_paragraph("")
        
    doc.save(DOCX_OUTPUT)
    print(f"Generated Word Doc: {DOCX_OUTPUT}")

if __name__ == '__main__':
    data = fetch_data()
    generate_pdf(data)
    generate_excel(data)
    generate_docx(data)
    print("All documents generated successfully with 100% clean formatting!")
